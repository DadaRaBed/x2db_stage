from pathlib import Path
import re
import pandas as pd
import sqlite3
import webview
import json
import os
import gc
import sys
import unicodedata
import time
from typing import Optional, Dict, Any, List
from datetime import datetime

from services.auth_service import AuthService
from services.database_service import DatabaseService
from services.excel_service import ExcelService

from repositories.system_database import initialize_database, user_count


# ============================================================
# DETECTION DU MODE D'EXECUTION (DEV vs EXE PyInstaller)
# ============================================================
def get_app_data_dir() -> Path:
    """
    Retourne le dossier ou stocker les donnees utilisateur (persistant).

    - Mode EXE PyInstaller : %APPDATA%/DataManager/data (Windows)
                             ~/.local/share/DataManager/data (Linux/macOS)
    - Mode DEV             : <projet>/data

    Le dossier est cree automatiquement s'il n'existe pas.
    """
    if getattr(sys, "frozen", False):
        # Mode EXE : utiliser un dossier PERSISTANT
        if sys.platform == "win32":
            base = Path(os.environ.get("APPDATA", os.path.expanduser("~")))
        else:
            base = Path(os.path.expanduser("~/.local/share"))
        app_dir = base / "DataManager"
    else:
        # Mode DEV : dossier du projet
        app_dir = Path(__file__).resolve().parent

    data_dir = app_dir / "data"
    data_dir.mkdir(parents=True, exist_ok=True)
    return data_dir


def get_resource_path(relative_path: str) -> Path:
    """
    Retourne le chemin vers une ressource embarquee (web/index.html...).

    - Mode EXE : _MEIPASS (dossier temporaire PyInstaller) ou dossier de l'exe
    - Mode DEV : dossier du projet
    """
    if getattr(sys, "frozen", False):
        base = Path(getattr(sys, "_MEIPASS", Path(sys.executable).parent))
    else:
        base = Path(__file__).resolve().parent
    return base / relative_path


# ============================================================
# CONSTANTES GLOBALES
# ============================================================
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = get_app_data_dir()
INDEX_FILE = get_resource_path("web/index.html")

_APP_WINDOW = None


class Api:
    def __init__(self):
        self.current_user = None
        self._database_service = DatabaseService()
        self._auth_service = AuthService()
        self._excel_service = ExcelService()
        self._active_db_path = None
        self._is_loading = False
        self._last_db_path = None
        self._loading_start_time = None

    # ============================================================
    # CONNEXION SQLITE COMPATIBLE MULTI-THREAD
    # ============================================================
    def _connect_db(self, db_path: str, row_factory=None):
        """
        Ouvre une connexion SQLite compatible multi-thread.

        pywebview execute les appels API dans un thread different du thread
        principal. Sans check_same_thread=False, SQLite leve l'erreur :
        'SQLite objects created in a thread can only be used in that same thread'.
        """
        conn = sqlite3.connect(db_path, check_same_thread=False, timeout=30.0)
        try:
            conn.execute("PRAGMA journal_mode=WAL;")
            conn.execute("PRAGMA synchronous=NORMAL;")
        except Exception:
            pass
        if row_factory is not None:
            conn.row_factory = row_factory
        return conn

    def _safe_close_connection(self, conn):
        """Ferme proprement une connexion SQLite avec checkpoint WAL."""
        if conn is None:
            return
        try:
            conn.execute("PRAGMA wal_checkpoint(TRUNCATE);")
        except Exception:
            pass
        try:
            conn.close()
        except Exception:
            pass

    def _release_resources(self, delay: float = 0.25):
        """Force le GC et laisse SQLite liberer les fichiers WAL/SHM."""
        try:
            gc.collect()
        except Exception:
            pass
        time.sleep(delay)

    # ============================================================
    # AUTHENTIFICATION
    # ============================================================
    def get_auth_status(self):
        try:
            return {
                "success": True,
                "first_user_exists": user_count() > 0,
                "authenticated": self.current_user is not None,
                "user": self.current_user,
            }
        except Exception as e:
            return {
                "success": False,
                "first_user_exists": False,
                "authenticated": False,
                "user": None,
                "message": f"Impossible de verifier le statut d'authentification : {e}",
            }

    def create_first_user(self, pseudo: str, password: str):
        try:
            return self._auth_service.create_first_user(pseudo, password)
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de creer l'utilisateur : {e}",
            }

    def login(self, pseudo: str, password: str):
        try:
            result = self._auth_service.login(pseudo, password)
            if result.get("success"):
                self.current_user = result.get("user")
            return result
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de se connecter : {e}",
            }

    def logout(self):
        try:
            if self._active_db_path is not None:
                return {
                    "success": False,
                    "message": "Une base de donnees est encore ouverte. Veuillez cliquer sur le bouton 'Terminer' pour fermer la base avant de vous deconnecter.",
                    "error_type": "database_still_open",
                    "active_db_path": self._active_db_path,
                }

            if self._is_loading:
                elapsed = time.time() - (self._loading_start_time or time.time())
                if elapsed < 5:
                    return {
                        "success": False,
                        "message": "Une operation est en cours. Veuillez patienter avant de vous deconnecter.",
                        "error_type": "operation_in_progress",
                    }
                else:
                    self._is_loading = False
                    self._loading_start_time = None

            self.current_user = None
            try:
                self._database_service.close_database()
            except Exception:
                pass
            self._active_db_path = None
            self._last_db_path = None
            self._is_loading = False
            self._loading_start_time = None

            self._release_resources(0.15)

            return {
                "success": True,
                "message": "Deconnexion reussie.",
            }
        except Exception:
            self.current_user = None
            self._active_db_path = None
            self._last_db_path = None
            self._is_loading = False
            return {
                "success": True,
                "message": "Deconnexion reussie.",
            }

    # ============================================================
    # TERMINER LA BASE DE DONNEES
    # ============================================================
    def terminate_database(self):
        try:
            if self._is_loading:
                elapsed = time.time() - (self._loading_start_time or time.time())
                if elapsed < 5:
                    return {
                        "success": False,
                        "message": "Une operation est en cours. Veuillez patienter quelques secondes.",
                        "error_type": "operation_in_progress",
                    }
                else:
                    self._is_loading = False
                    self._loading_start_time = None

            if self._active_db_path is None:
                return {
                    "success": True,
                    "message": "Aucune base de donnees n'est actuellement ouverte.",
                    "already_closed": True,
                }

            db_name = os.path.basename(self._active_db_path)

            try:
                self._database_service.close_database()
            except Exception as e:
                print(f"[INFO] Fermeture base (thread different) : {e}")

            self._active_db_path = None
            self._last_db_path = None
            self._is_loading = False
            self._loading_start_time = None

            self._release_resources(0.3)

            return {
                "success": True,
                "message": f"Base de donnees '{db_name}' terminee avec succes. Vous pouvez maintenant vous deconnecter.",
                "closed_db": db_name,
            }

        except Exception:
            self._active_db_path = None
            self._last_db_path = None
            self._is_loading = False
            self._loading_start_time = None
            return {
                "success": True,
                "message": "Base de donnees terminee.",
            }

    def quit_app(self):
        global _APP_WINDOW
        try:
            if _APP_WINDOW:
                _APP_WINDOW.destroy()
            os._exit(0)
        except Exception:
            os._exit(0)

    # ============================================================
    # ACTIVITES
    # ============================================================
    def get_activities(self, limit: int = 5):
        try:
            log_file = DATA_DIR / "activity_log.json"
            if not log_file.exists():
                return {"success": True, "activities": []}

            with open(log_file, "r", encoding="utf-8") as f:
                activities = json.load(f)

            if not isinstance(activities, list):
                return {"success": True, "activities": []}

            activities.sort(key=lambda x: x.get("date", ""), reverse=True)
            return {"success": True, "activities": activities[:limit]}
        except Exception as e:
            return {"success": False, "message": str(e), "activities": []}

    def log_activity(self, action_text: str):
        try:
            log_file = DATA_DIR / "activity_log.json"
            log_file.parent.mkdir(parents=True, exist_ok=True)

            activities = []
            if log_file.exists():
                with open(log_file, "r", encoding="utf-8") as f:
                    try:
                        activities = json.load(f)
                    except json.JSONDecodeError:
                        activities = []

            if not isinstance(activities, list):
                activities = []

            activities.append({
                "text": action_text,
                "date": time.strftime("%Y-%m-%d %H:%M:%S"),
                "user": self.current_user.get("pseudo") if self.current_user else "Inconnu",
            })

            if len(activities) > 100:
                activities = activities[-100:]

            with open(log_file, "w", encoding="utf-8") as f:
                json.dump(activities, f, ensure_ascii=False, indent=2)

            return {"success": True}
        except Exception as e:
            return {"success": False, "message": str(e)}

    # ============================================================
    # SUPPRESSION DE BASE DE DONNEES
    # ============================================================
    def delete_database(self, db_path: str):
        """
        Supprime definitivement une base de donnees (.db) du dossier data/.

        Protections :
        - system.db ne peut JAMAIS etre supprime
        - La base actuellement ouverte est fermee automatiquement
        - Fichiers annexes WAL/SHM supprimes aussi
        """
        try:
            if not db_path or not str(db_path).strip():
                return {"success": False, "message": "Le chemin est vide."}

            db_file = Path(db_path).resolve()

            # Securite : reste dans le dossier DATA_DIR
            try:
                db_file.relative_to(DATA_DIR.resolve())
            except ValueError:
                return {
                    "success": False,
                    "message": "Suppression refusee : la base doit se trouver dans le dossier data/.",
                }

            if not db_file.exists():
                return {"success": False, "message": "La base de donnees n'existe pas."}

            if db_file.name.lower() == "system.db":
                return {
                    "success": False,
                    "message": "Suppression interdite : 'system.db' contient les informations des utilisateurs.",
                }

            db_name = db_file.name

            if self._active_db_path and os.path.abspath(self._active_db_path) == str(db_file):
                try:
                    self._database_service.close_database()
                except Exception:
                    pass
                self._active_db_path = None
                self._last_db_path = None
                self._release_resources(0.2)

            time.sleep(0.25)

            deleted_files = []
            try:
                if db_file.exists():
                    os.remove(str(db_file))
                    deleted_files.append(db_name)
            except Exception as e:
                return {
                    "success": False,
                    "message": f"Impossible de supprimer '{db_name}' : {e}",
                }

            for suffix in ("-wal", "-shm", "-journal"):
                annex = Path(str(db_file) + suffix)
                if annex.exists():
                    try:
                        os.remove(str(annex))
                        deleted_files.append(annex.name)
                    except Exception:
                        pass

            return {
                "success": True,
                "message": f"Base de donnees '{db_name}' supprimee definitivement.",
                "deleted": deleted_files,
                "db_name": db_name,
            }

        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"success": False, "message": f"Erreur lors de la suppression : {e}"}

    # ============================================================
    # SUPPRESSION DE TABLE
    # ============================================================
    def delete_table(self, table_name: str, file_path: str = None):
        """Supprime une table dans la base active (sauf tables systeme)."""
        try:
            if not table_name or not str(table_name).strip():
                return {"success": False, "message": "Le nom de la table est requis."}

            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            safe_table = str(table_name).replace('"', '""')

            conn = self._connect_db(db_path)
            try:
                cursor = conn.cursor()

                cursor.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name = ?",
                    (table_name,),
                )
                if not cursor.fetchone():
                    return {
                        "success": False,
                        "message": f"La table '{table_name}' n'existe pas.",
                    }

                if table_name.lower() in ("sqlite_sequence", "sqlite_master"):
                    return {
                        "success": False,
                        "message": "Suppression interdite : table systeme SQLite.",
                    }

                cursor.execute(f'DROP TABLE IF EXISTS "{safe_table}"')
                conn.commit()
            finally:
                self._safe_close_connection(conn)

            return {
                "success": True,
                "message": f"Table '{table_name}' supprimee avec succes.",
                "table_name": table_name,
            }

        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"success": False, "message": f"Erreur lors de la suppression : {e}"}

    # ============================================================
    # FORMATAGE EXCEL PROFESSIONNEL
    # ============================================================
    def _apply_professional_formatting(self, writer, sheet_name: str, df: pd.DataFrame):
        """Applique une mise en forme professionnelle a une feuille Excel."""
        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            worksheet = writer.sheets[sheet_name]
            max_row = len(df) + 1
            max_col = len(df.columns)

            if max_col == 0:
                return

            worksheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(
                start_color="4F46E5", end_color="4F46E5", fill_type="solid"
            )
            header_alignment = Alignment(
                horizontal="center", vertical="center", wrap_text=True
            )
            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            for col_idx in range(1, max_col + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            alt_fill_1 = PatternFill(
                start_color="FFFFFF", end_color="FFFFFF", fill_type="solid"
            )
            alt_fill_2 = PatternFill(
                start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
            )
            cell_alignment = Alignment(vertical="center", wrap_text=False)

            for row_idx in range(2, max_row + 1):
                fill = alt_fill_1 if row_idx % 2 == 0 else alt_fill_2
                for col_idx in range(1, max_col + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.fill = fill
                    cell.alignment = cell_alignment

            worksheet.freeze_panes = "A2"

            for col_idx, column in enumerate(df.columns, 1):
                col_letter = get_column_letter(col_idx)
                max_length = len(str(column))
                for row_idx in range(2, min(len(df) + 2, 500)):
                    cell_value = worksheet.cell(row=row_idx, column=col_idx).value
                    if cell_value is not None:
                        cell_len = len(str(cell_value))
                        if cell_len > max_length:
                            max_length = cell_len
                adjusted_width = min(max_length + 3, 50)
                worksheet.column_dimensions[col_letter].width = adjusted_width

            worksheet.row_dimensions[1].height = 25

        except Exception as e:
            print(f"[WARN] Erreur formatage Excel pour '{sheet_name}' : {e}")

    # ============================================================
    # CREATION DE BASE DE DONNEES
    # ============================================================
    def create_new_database(self, db_name: str, table_name: str):
        try:
            if not db_name or not db_name.strip():
                return {"success": False, "message": "Le nom de la base est requis."}

            if not table_name or not table_name.strip():
                return {"success": False, "message": "Le nom de la table est requis."}

            data_dir = DATA_DIR
            data_dir.mkdir(parents=True, exist_ok=True)

            safe_db_name = self._clean_ascii(db_name) or "nouvelle_base"
            safe_table_name = self._clean_ascii(table_name) or "nouvelle_table"

            db_filename = f"{safe_db_name}.db"
            db_path = data_dir / db_filename

            if os.path.exists(str(db_path)):
                return {
                    "success": False,
                    "message": f"Une base de donnees nommee '{db_filename}' existe deja. Veuillez choisir un autre nom.",
                }

            conn = self._connect_db(str(db_path))
            try:
                cursor = conn.cursor()
                cursor.execute(f'''
                    CREATE TABLE IF NOT EXISTS "{safe_table_name}" (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        cin TEXT,
                        nom TEXT,
                        commune TEXT,
                        fkt TEXT,
                        district TEXT,
                        region TEXT,
                        filiation_menage TEXT,
                        pole_de_developpement TEXT,
                        filieres TEXT,
                        opr TEXT,
                        h_f TEXT,
                        categorisation_eaf TEXT,
                        variete TEXT,
                        observation TEXT
                    )
                ''')
                conn.commit()
            finally:
                self._safe_close_connection(conn)

            return {
                "success": True,
                "message": f"Base de donnees '{db_filename}' creee avec succes.",
                "db_path": str(db_path),
                "db_name": db_filename,
                "table_name": safe_table_name,
            }

        except Exception as e:
            return {"success": False, "message": f"Erreur lors de la creation : {e}"}

    def create_database_from_excel(self, file_path: str, db_name: str = None):
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le fichier Excel est requis."}

            excel_path = Path(file_path)
            if not os.path.exists(str(excel_path)):
                return {"success": False, "message": f"Fichier introuvable : {file_path}"}

            data_dir = DATA_DIR
            data_dir.mkdir(parents=True, exist_ok=True)

            if db_name:
                safe_db_name = self._clean_ascii(db_name) or "database"
            else:
                safe_db_name = self._clean_ascii(excel_path.stem) or "database"

            db_filename = f"{safe_db_name}.db"
            db_path = data_dir / db_filename

            all_sheets = pd.read_excel(excel_path, sheet_name=None)
            conn = self._connect_db(str(db_path))
            tables_created = []
            total_rows = 0

            try:
                for current_sheet, df in all_sheets.items():
                    clean_table_name = self._clean_ascii(str(current_sheet)) or "table"

                    try:
                        df = df.dropna(how="all")

                        if any(str(col).lower().startswith("unnamed") for col in df.columns):
                            if len(df) > 0:
                                new_headers = df.iloc[0].fillna("colonne_inconnue").astype(str).tolist()
                                cleaned_headers = []
                                seen = {}
                                for h in new_headers:
                                    h_clean = self._clean_ascii(h) or "col"
                                    if h_clean in seen:
                                        seen[h_clean] += 1
                                        h_clean = f"{h_clean}_{seen[h_clean]}"
                                    else:
                                        seen[h_clean] = 0
                                    cleaned_headers.append(h_clean)
                                df.columns = cleaned_headers
                                df = df.drop(df.index[0])

                        for col in df.columns:
                            if pd.api.types.is_numeric_dtype(df[col]):
                                df[col] = df[col].fillna(0)
                            elif pd.api.types.is_datetime64_any_dtype(df[col]):
                                df[col] = pd.to_datetime(df[col]).dt.date
                                df[col] = df[col].fillna(pd.Timestamp.now().date())
                            else:
                                df[col] = df[col].fillna("Non specifie")

                        df.columns = [self._clean_ascii(str(col)) or "col" for col in df.columns]
                        df.to_sql(clean_table_name, conn, if_exists="replace", index=False)
                        tables_created.append(clean_table_name)
                        total_rows += len(df)
                    except Exception as sheet_err:
                        print(f"[ERREUR IMPORT] Feuille '{current_sheet}' : {sheet_err}")
                        continue
            finally:
                self._safe_close_connection(conn)

            return {
                "success": True,
                "message": f"Base de donnees '{db_filename}' creee avec {len(tables_created)} table(s).",
                "db_path": str(db_path),
                "db_name": db_filename,
                "tables": tables_created,
                "total_rows": total_rows,
            }

        except Exception as e:
            return {"success": False, "message": f"Erreur lors de la creation : {e}"}

    # ============================================================
    # EXPORT EXCEL
    # ============================================================
    def export_database_to_excel_from_path(self, db_path: str, output_excel_path: str):
        """Export professionnel : chaque table devient une feuille."""
        try:
            if not db_path or not os.path.exists(db_path):
                return {"success": False, "message": "Base de donnees introuvable."}

            conn = self._connect_db(db_path)
            try:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
                )
                tables = [
                    row[0] for row in cursor.fetchall() if row[0] != "sqlite_sequence"
                ]

                if not tables:
                    return {"success": False, "message": "Aucune table dans cette base."}

                tables_ordered = [t for t in tables if t != "listes_meres"]
                if "listes_meres" in tables:
                    tables_ordered.append("listes_meres")

                with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
                    for table in tables_ordered:
                        df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)

                        if "id" in df.columns:
                            df = df.drop(columns=["id"])

                        for col in df.columns:
                            if pd.api.types.is_datetime64_any_dtype(df[col]):
                                df[col] = df[col].dt.date

                        sheet_name = table[:31] if len(table) <= 31 else table[:31]
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        self._apply_professional_formatting(writer, sheet_name, df)
            finally:
                self._safe_close_connection(conn)

            return {
                "success": True,
                "message": f"Exportation professionnelle reussie vers {output_excel_path}",
            }
        except Exception as e:
            return {"success": False, "message": str(e)}

    def export_database_to_excel(self, output_excel_path: str, file_path: str = None):
        """Export professionnel de la base active."""
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            conn = self._connect_db(db_path)
            try:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
                )
                tables = [
                    row[0] for row in cursor.fetchall() if row[0] != "sqlite_sequence"
                ]

                tables_ordered = [t for t in tables if t != "listes_meres"]
                if "listes_meres" in tables:
                    tables_ordered.append("listes_meres")

                with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
                    for table in tables_ordered:
                        df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)

                        if "id" in df.columns:
                            df = df.drop(columns=["id"])

                        for col in df.columns:
                            if pd.api.types.is_datetime64_any_dtype(df[col]):
                                df[col] = df[col].dt.date

                        sheet_name = table[:31] if len(table) <= 31 else table[:31]
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        self._apply_professional_formatting(writer, sheet_name, df)
            finally:
                self._safe_close_connection(conn)

            return {
                "success": True,
                "message": f"Exportation reussie vers {output_excel_path}",
            }
        except Exception as e:
            return {"success": False, "message": str(e)}

    # ============================================================
    # CREATION DES LISTES MERES  (SANS DEDOUBLONNAGE)
    # ============================================================
    def create_master_list(self, file_path: str = None):
        """
        Cree une table 'listes_meres' contenant TOUTES les personnes de TOUTES
        les tables, SANS aucun dedoublonnage.

        Ajoute deux colonnes de tracabilite :
        - source_table  : nom de la table d'origine
        - ligne_origine : rowid de la ligne dans sa table d'origine

        L'utilisateur peut ensuite utiliser l'outil 'Trouver les doublons'
        pour identifier et supprimer manuellement les doublons CIN+NOM.

        Seules les lignes dont le nom est vide sont ignorees.
        """
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            # Fermer la connexion eventuellement ouverte par DatabaseService
            try:
                self._database_service.close_database()
            except Exception:
                pass

            self._release_resources(0.15)

            conn = self._connect_db(db_path)
            try:
                cursor = conn.cursor()

                cursor.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
                )
                tables = [
                    row[0]
                    for row in cursor.fetchall()
                    if row[0] not in ("sqlite_sequence", "listes_meres")
                ]

                if not tables:
                    return {"success": False, "message": "Aucune table dans cette base."}

                tables = sorted(tables)

                all_persons = []
                total_lignes_lues = 0
                total_noms_vides_ignores = 0

                def norm(v):
                    """Normalise une valeur en chaine propre (sans 'nan'/None)."""
                    if v is None:
                        return ""
                    s = str(v).strip()
                    if s.upper() in ("NAN", "NONE", "NULL"):
                        return ""
                    return s

                for table in tables:
                    try:
                        # Lire avec le rowid d'origine pour tracabilite
                        df = pd.read_sql_query(
                            f'SELECT rowid AS __ligne_origine__, * FROM "{table}"',
                            conn,
                        )
                        if df.empty:
                            continue

                        total_lignes_lues += len(df)

                        # --- Detection des colonnes ---
                        cin_col = nom_col = commune_col = fkt_col = annee_col = hf_col = None
                        region_col = district_col = filieres_col = None
                        categorisation_col = variete_col = opr_col = None
                        filiation_col = pole_col = None

                        for c in df.columns:
                            cl = c.lower()
                            if "cin" in cl and not cin_col:
                                cin_col = c
                            if ("nom" in cl or "prenom" in cl) and not nom_col:
                                nom_col = c
                            if "commune" in cl and not commune_col:
                                commune_col = c
                            if ("fkt" in cl or "fokontany" in cl) and not fkt_col:
                                fkt_col = c
                            if ("annee" in cl and "naissance" in cl) and not annee_col:
                                annee_col = c
                            if (cl in ("h_f", "hf", "sexe")) and not hf_col:
                                hf_col = c
                            if cl == "region":
                                region_col = c
                            elif cl == "district":
                                district_col = c
                            elif cl in ("filieres", "filiere"):
                                filieres_col = c
                            elif "categorisation" in cl or "categoris" in cl:
                                categorisation_col = c
                            elif cl == "variete":
                                variete_col = c
                            elif cl == "opr":
                                opr_col = c
                            elif "filiation" in cl:
                                filiation_col = c
                            elif "pole" in cl or "podev" in cl:
                                pole_col = c

                        # --- Lecture ligne par ligne SANS dedoublonnage ---
                        for idx, row in df.iterrows():
                            nom_val = norm(row.get(nom_col, "")) if nom_col else ""

                            # Ignorer seulement les lignes sans nom
                            if not nom_val:
                                total_noms_vides_ignores += 1
                                continue

                            person = {
                                "source_table": table,
                                "ligne_origine": int(row["__ligne_origine__"]),
                                "region": norm(row.get(region_col, "")) if region_col else "",
                                "district": norm(row.get(district_col, "")) if district_col else "",
                                "commune": norm(row.get(commune_col, "")) if commune_col else "",
                                "fkt": norm(row.get(fkt_col, "")) if fkt_col else "",
                                "nom_et_prenoms": nom_val,
                                "h_f": norm(row.get(hf_col, "")) if hf_col else "",
                                "filieres": norm(row.get(filieres_col, "")) if filieres_col else "",
                                "cin": norm(row.get(cin_col, "")) if cin_col else "",
                                "annee_de_naissance": norm(row.get(annee_col, "")) if annee_col else "",
                                "categorisation_eaf": norm(row.get(categorisation_col, "")) if categorisation_col else "",
                                "variete": norm(row.get(variete_col, "")) if variete_col else "",
                                "opr": norm(row.get(opr_col, "")) if opr_col else "",
                                "filiation_menage": norm(row.get(filiation_col, "")) if filiation_col else "",
                                "pole_de_developpement": norm(row.get(pole_col, "")) if pole_col else "",
                            }

                            all_persons.append(person)

                    except Exception as e:
                        print(f"[WARN] Erreur lecture table '{table}' : {e}")
                        continue

                if not all_persons:
                    return {
                        "success": False,
                        "message": "Aucune personne valide trouvee.",
                    }

                # --- Ecriture de la table listes_meres ---
                cursor.execute('DROP TABLE IF EXISTS "listes_meres"')
                cursor.execute('''
                    CREATE TABLE "listes_meres" (
                        id INTEGER PRIMARY KEY AUTOINCREMENT,
                        source_table TEXT,
                        ligne_origine INTEGER,
                        region TEXT,
                        district TEXT,
                        commune TEXT,
                        fkt TEXT,
                        nom_et_prenoms TEXT,
                        h_f TEXT,
                        filieres TEXT,
                        cin TEXT,
                        annee_de_naissance TEXT,
                        categorisation_eaf TEXT,
                        variete TEXT,
                        opr TEXT,
                        filiation_menage TEXT,
                        pole_de_developpement TEXT
                    )
                ''')

                for p in all_persons:
                    cursor.execute('''
                        INSERT INTO "listes_meres"
                        (source_table, ligne_origine, region, district, commune,
                         fkt, nom_et_prenoms, h_f, filieres, cin,
                         annee_de_naissance, categorisation_eaf, variete, opr,
                         filiation_menage, pole_de_developpement)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        p.get("source_table", ""),
                        p.get("ligne_origine", 0),
                        p.get("region", ""),
                        p.get("district", ""),
                        p.get("commune", ""),
                        p.get("fkt", ""),
                        p.get("nom_et_prenoms", ""),
                        p.get("h_f", ""),
                        p.get("filieres", ""),
                        p.get("cin", ""),
                        p.get("annee_de_naissance", ""),
                        p.get("categorisation_eaf", ""),
                        p.get("variete", ""),
                        p.get("opr", ""),
                        p.get("filiation_menage", ""),
                        p.get("pole_de_developpement", ""),
                    ))

                conn.commit()
            finally:
                self._safe_close_connection(conn)

            S = len(all_persons)
            N = total_lignes_lues

            return {
                "success": True,
                "message": (
                    f"Liste mere creee avec {S} personne(s) (TOUTES les lignes brutes, "
                    f"AUCUN dedoublonnage). "
                    f"[{N} lignes lues, {total_noms_vides_ignores} noms vides ignores] "
                    f"Utilisez maintenant l'outil 'Trouver les doublons' pour "
                    f"identifier et supprimer manuellement les doublons CIN+NOM."
                ),
                "total_persons": S,
                "tables_scanned": len(tables),
                "total_lignes_lues": N,
                "total_noms_vides_ignores": total_noms_vides_ignores,
            }

        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"success": False, "message": f"Erreur lors de la creation : {e}"}

    # ============================================================
    # REQUETES STATISTIQUES
    # ============================================================
    def execute_statistical_query(self, query_type: str, params: Dict[str, Any] = None, file_path: str = None):
        """Execute des requetes statistiques predefinies."""
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "data": []}

            if params is None:
                params = {}

            conn = self._connect_db(db_path)
            conn.row_factory = sqlite3.Row

            try:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
                )
                tables = [
                    row[0] for row in cursor.fetchall() if row[0] not in ("sqlite_sequence",)
                ]

                data = []

                if query_type == "femmes_par_tranche_age":
                    age_min = int(params.get("age_min", 0))
                    age_max = int(params.get("age_max", 100))
                    buckets = {
                        "0-14": 0, "15-24": 0, "25-34": 0,
                        "35-44": 0, "45-54": 0, "55-64": 0, "65+": 0,
                    }
                    current_year = datetime.now().year

                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            hf_col = None
                            annee_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if hf_col is None and (cl == "h_f" or cl == "hf" or cl == "sexe"):
                                    hf_col = c
                                if annee_col is None and ("annee" in cl and "naissance" in cl):
                                    annee_col = c

                            if not hf_col or not annee_col:
                                continue

                            for idx, row in df.iterrows():
                                hf = str(row[hf_col] or "").strip().upper()
                                is_female = hf in ("F", "FEMME", "FEMININ", "FEMALE", "2")
                                if not is_female:
                                    continue

                                try:
                                    annee = int(float(str(row[annee_col] or 0)))
                                    if annee < 1900 or annee > current_year:
                                        continue
                                    age = current_year - annee
                                except Exception:
                                    continue

                                if age < age_min or age > age_max:
                                    continue

                                if age <= 14:
                                    buckets["0-14"] += 1
                                elif age <= 24:
                                    buckets["15-24"] += 1
                                elif age <= 34:
                                    buckets["25-34"] += 1
                                elif age <= 44:
                                    buckets["35-44"] += 1
                                elif age <= 54:
                                    buckets["45-54"] += 1
                                elif age <= 64:
                                    buckets["55-64"] += 1
                                else:
                                    buckets["65+"] += 1
                        except Exception as e:
                            print(f"[WARN] {table} : {e}")
                            continue

                    for tranche, count in buckets.items():
                        data.append({"tranche_age": tranche, "nombre_femmes": count})

                elif query_type == "personnes_par_lieu":
                    lieu_type = params.get("lieu_type", "commune")
                    counts = {}
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            lieu_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if lieu_type in cl:
                                    lieu_col = c
                                    break

                            if not lieu_col:
                                continue

                            for val in df[lieu_col]:
                                v = str(val or "").strip()
                                if v and v.lower() != "nan" and v.lower() != "non specifie":
                                    counts[v] = counts.get(v, 0) + 1
                        except Exception:
                            continue

                    sorted_counts = sorted(counts.items(), key=lambda x: -x[1])
                    for lieu, count in sorted_counts:
                        data.append({"lieu": lieu, "nombre_personnes": count})

                elif query_type == "superficie_par_personne":
                    superficies = {}
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            nom_col = None
                            cin_col = None
                            sup_col = None

                            for c in df.columns:
                                cl = c.lower()
                                if "nom" in cl and not nom_col:
                                    nom_col = c
                                if "cin" in cl and not cin_col:
                                    cin_col = c
                                if (
                                    "superficie" in cl
                                    or "superf" in cl
                                    or "surface" in cl
                                    or cl == "ha"
                                ):
                                    sup_col = c

                            if not sup_col:
                                continue

                            for idx, row in df.iterrows():
                                key = (
                                    str(row.get(nom_col, "") or "").strip()
                                    + "|"
                                    + str(row.get(cin_col, "") or "").strip()
                                )
                                try:
                                    sup = float(str(row[sup_col] or 0).replace(",", "."))
                                except Exception:
                                    sup = 0
                                superficies[key] = superficies.get(key, 0) + sup
                        except Exception:
                            continue

                    sorted_sup = sorted(superficies.items(), key=lambda x: -x[1])
                    for key, sup in sorted_sup[:100]:
                        parts = key.split("|")
                        data.append({
                            "nom": parts[0] if len(parts) > 0 else "",
                            "cin": parts[1] if len(parts) > 1 else "",
                            "superficie_totale": round(sup, 2),
                        })

                elif query_type == "hommes_femmes":
                    counts = {"H": 0, "F": 0, "Autre": 0}
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            hf_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if cl in ("h_f", "hf", "sexe"):
                                    hf_col = c
                                    break

                            if not hf_col:
                                continue

                            for val in df[hf_col]:
                                v = str(val or "").strip().upper()
                                if v in ("H", "HOMME", "MASCULIN", "MALE", "1"):
                                    counts["H"] += 1
                                elif v in ("F", "FEMME", "FEMININ", "FEMALE", "2"):
                                    counts["F"] += 1
                                elif v:
                                    counts["Autre"] += 1
                        except Exception:
                            continue

                    for sexe, count in counts.items():
                        data.append({"sexe": sexe, "nombre": count})

                elif query_type == "personnes_par_filiere":
                    counts = {}
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            fil_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if "filiere" in cl:
                                    fil_col = c
                                    break

                            if not fil_col:
                                continue

                            for val in df[fil_col]:
                                v = str(val or "").strip()
                                if v and v.lower() != "nan":
                                    counts[v] = counts.get(v, 0) + 1
                        except Exception:
                            continue

                    sorted_c = sorted(counts.items(), key=lambda x: -x[1])
                    for fil, count in sorted_c:
                        data.append({"filiere": fil, "nombre_personnes": count})

                elif query_type == "personnes_par_commune":
                    counts = {}
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            comm_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if "commune" in cl:
                                    comm_col = c
                                    break

                            if not comm_col:
                                continue

                            for val in df[comm_col]:
                                v = str(val or "").strip()
                                if v and v.lower() != "nan":
                                    counts[v] = counts.get(v, 0) + 1
                        except Exception:
                            continue

                    sorted_c = sorted(counts.items(), key=lambda x: -x[1])
                    for comm, count in sorted_c:
                        data.append({"commune": comm, "nombre_personnes": count})

                elif query_type == "personnes_par_categorisation":
                    counts = {}
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            cat_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if "categorisation" in cl or "categoris" in cl:
                                    cat_col = c
                                    break

                            if not cat_col:
                                continue

                            for val in df[cat_col]:
                                v = str(val or "").strip()
                                if v and v.lower() != "nan":
                                    counts[v] = counts.get(v, 0) + 1
                        except Exception:
                            continue

                    sorted_c = sorted(counts.items(), key=lambda x: -x[1])
                    for cat, count in sorted_c:
                        data.append({"categorisation": cat, "nombre_personnes": count})

                elif query_type == "age_moyen":
                    current_year = datetime.now().year
                    total_age = 0
                    count = 0

                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            annee_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if "annee" in cl and "naissance" in cl:
                                    annee_col = c
                                    break

                            if not annee_col:
                                continue

                            for val in df[annee_col]:
                                try:
                                    annee = int(float(str(val)))
                                    if 1900 < annee <= current_year:
                                        total_age += current_year - annee
                                        count += 1
                                except Exception:
                                    continue
                        except Exception:
                            continue

                    avg = round(total_age / count, 1) if count > 0 else 0
                    data.append({"age_moyen": avg, "total_personnes": count})

                elif query_type == "superficie_totale":
                    total = 0
                    for table in tables:
                        try:
                            df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                            if df.empty:
                                continue

                            sup_col = None
                            for c in df.columns:
                                cl = c.lower()
                                if (
                                    "superficie" in cl
                                    or "superf" in cl
                                    or "surface" in cl
                                ):
                                    sup_col = c
                                    break

                            if not sup_col:
                                continue

                            for val in df[sup_col]:
                                try:
                                    total += float(str(val).replace(",", "."))
                                except Exception:
                                    continue
                        except Exception:
                            continue

                    data.append({"superficie_totale": round(total, 2)})

                else:
                    return {
                        "success": False,
                        "message": f"Type de requete inconnu : {query_type}",
                        "data": [],
                    }

                return {"success": True, "data": data, "query_type": query_type}
            finally:
                self._safe_close_connection(conn)

        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"success": False, "message": str(e), "data": []}

    # ============================================================
    # GESTION DES BASES
    # ============================================================
    def get_database_info(self):
        try:
            return self._database_service.get_database_info()
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de recuperer les informations : {e}",
            }

    def get_data_directory_databases(self):
        try:
            data_dir = DATA_DIR
            if not data_dir.exists() or not data_dir.is_dir():
                return {"success": True, "databases": []}

            databases = []
            for file_path in data_dir.glob("*.db"):
                if file_path.is_file() and file_path.name.lower() != "system.db":
                    size_kb = round(file_path.stat().st_size / 1024, 2)
                    databases.append({
                        "name": file_path.name,
                        "path": str(file_path.resolve()),
                        "size_kb": size_kb,
                    })

            return {"success": True, "databases": databases}
        except Exception as e:
            return {"success": False, "message": str(e), "databases": []}

    # ============================================================
    # OUVRIR UNE BASE (version renforcee anti-thread-zombie)
    # ============================================================
    def open_database(self, path: str):
        try:
            if self._is_loading:
                return {"success": False, "message": "Une operation est deja en cours."}

            if not path:
                return {"success": False, "message": "Le chemin est vide."}

            db_path = Path(path)
            if not db_path.is_absolute():
                db_path = DATA_DIR / db_path.name

            resolved_path = os.path.abspath(str(db_path))
            if not os.path.exists(resolved_path):
                return {
                    "success": False,
                    "message": f"La base de donnees n'existe pas : {resolved_path}",
                }

            # Si on reouvre la meme base qui est deja ouverte, on la ferme
            if self._active_db_path == resolved_path:
                try:
                    self._database_service.close_database()
                except Exception:
                    pass
                self._active_db_path = None
                self._last_db_path = None
                self._release_resources(0.2)

            self._is_loading = True
            self._loading_start_time = time.time()

            # Fermer toute base precedemment ouverte
            if self._active_db_path and self._active_db_path != resolved_path:
                try:
                    self._database_service.close_database()
                except Exception:
                    pass
                self._active_db_path = None
                self._last_db_path = None
                self._release_resources(0.2)

            result = self._database_service.open_database(resolved_path)

            self._is_loading = False
            self._loading_start_time = None

            if result.get("success"):
                self._active_db_path = resolved_path
                self._last_db_path = resolved_path
                return result
            else:
                return result

        except Exception as e:
            self._is_loading = False
            self._loading_start_time = None
            import traceback
            traceback.print_exc()
            return {
                "success": False,
                "message": f"Impossible d'ouvrir la base : {e}",
            }

    def open_database_path(self, path: str):
        return self.open_database(path)

    def close_database(self):
        try:
            self._active_db_path = None
            self._last_db_path = None
            self._is_loading = False
            self._loading_start_time = None
            try:
                result = self._database_service.close_database()
            except Exception:
                result = {"success": True, "message": "Base fermee."}
            self._release_resources(0.15)
            return result
        except Exception:
            self._active_db_path = None
            self._last_db_path = None
            self._is_loading = False
            return {"success": True, "message": "Base fermee."}

    def _get_db_path(self, file_path: str = None) -> Optional[str]:
        if file_path:
            return file_path

        db_info = self._database_service.get_database_info()
        if db_info.get("success") and db_info.get("path"):
            return db_info["path"]

        return self._active_db_path

    # ============================================================
    # STRUCTURE ET DONNEES
    # ============================================================
    def get_database_structure_matrix(self, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "structure": {}}

            if not os.path.exists(db_path):
                return {
                    "success": False,
                    "message": f"La base n'existe pas : {db_path}",
                    "structure": {},
                }

            conn = self._connect_db(db_path)
            cursor = conn.cursor()
            cursor.execute(
                "SELECT name FROM sqlite_master WHERE type = 'table' AND name NOT LIKE 'sqlite_%' ORDER BY name"
            )
            tables = [row[0] for row in cursor.fetchall()]
            structure = {}

            for table_name in tables:
                escaped_table_name = table_name.replace('"', '""')
                cursor.execute(f'PRAGMA table_info("{escaped_table_name}")')
                columns = [column[1] for column in cursor.fetchall()]
                structure[table_name] = columns

            return {
                "success": True,
                "message": "Structure recuperee.",
                "structure": structure,
            }
        except Exception as error:
            return {
                "success": False,
                "message": f"Erreur : {str(error)}",
                "structure": {},
            }
        finally:
            self._safe_close_connection(conn)

    def get_database_table_names(self, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "tables": []}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()
            cursor.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
            )
            tables = [
                row[0] for row in cursor.fetchall() if row[0] != "sqlite_sequence"
            ]
            return {"success": True, "tables": tables}
        except Exception as e:
            return {"success": False, "message": str(e), "tables": []}
        finally:
            self._safe_close_connection(conn)

    def get_table_rows(self, table_name: str, file_path: str = None, limit: int = 1000):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "data": []}

            conn = self._connect_db(db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            cursor.execute(f'SELECT * FROM "{table_name}" LIMIT {limit}')
            rows = [dict(row) for row in cursor.fetchall()]
            return {"success": True, "data": rows}
        except Exception as e:
            return {"success": False, "message": str(e), "data": []}
        finally:
            self._safe_close_connection(conn)

    def get_table_columns(self, table_name: str, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "columns": []}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()
            escaped_table_name = table_name.replace('"', '""')
            cursor.execute(f'PRAGMA table_info("{escaped_table_name}")')
            columns = [col[1] for col in cursor.fetchall()]
            return {"success": True, "columns": columns}
        except Exception as e:
            return {"success": False, "message": str(e), "columns": []}
        finally:
            self._safe_close_connection(conn)

    def get_table_rows_filtered(
        self, table_name: str, columns: List[str], file_path: str = None
    ):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "data": []}

            conn = self._connect_db(db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()

            cols_sql = ", ".join([f'"{c}"' for c in columns]) if columns else "*"
            cursor.execute(f'SELECT {cols_sql} FROM "{table_name}" LIMIT 500')
            rows = [dict(row) for row in cursor.fetchall()]
            return {"success": True, "data": rows}
        except Exception as e:
            return {"success": False, "message": str(e), "data": []}
        finally:
            self._safe_close_connection(conn)

    def get_distinct_values(self, table_name: str, column: str, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "values": []}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()

            safe_table = table_name.replace('"', '""')
            safe_col = column.replace('"', '""')

            cursor.execute(
                f'SELECT DISTINCT "{safe_col}" FROM "{safe_table}" '
                f'WHERE "{safe_col}" IS NOT NULL AND "{safe_col}" != "" '
                f'ORDER BY "{safe_col}" LIMIT 200'
            )
            values = [row[0] for row in cursor.fetchall()]

            return {"success": True, "values": values}
        except Exception as e:
            return {"success": False, "message": str(e), "values": []}
        finally:
            self._safe_close_connection(conn)

    def search_in_table(self, table_name: str, search_term: str, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "data": []}

            conn = self._connect_db(db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()

            cursor.execute(f'PRAGMA table_info("{table_name}")')
            columns = [col[1] for col in cursor.fetchall()]

            conditions = " OR ".join([f'"{col}" LIKE ?' for col in columns])
            query = f'SELECT * FROM "{table_name}" WHERE {conditions}'
            params = [f"%{search_term}%" for _ in columns]

            cursor.execute(query, params)
            rows = [dict(row) for row in cursor.fetchall()]
            return {"success": True, "data": rows}
        except Exception as e:
            return {"success": False, "message": str(e), "data": []}
        finally:
            self._safe_close_connection(conn)

    # ============================================================
    # MODIFICATION DE LIGNE
    # ============================================================
    def update_table_row(
        self, table_name: str, row_id: int, column: str, value: str, file_path: str = None
    ):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()

            safe_table = table_name.replace('"', '""')
            safe_column = column.replace('"', '""')

            query = f'UPDATE "{safe_table}" SET "{safe_column}" = ? WHERE rowid = ?'
            cursor.execute(query, (value, row_id))
            conn.commit()

            if cursor.rowcount > 0:
                return {"success": True, "message": "Valeur modifiee avec succes."}
            else:
                return {"success": False, "message": "Aucune ligne modifiee."}
        except Exception as e:
            return {"success": False, "message": str(e)}
        finally:
            self._safe_close_connection(conn)

    def insert_table_row(
        self, table_name: str, values: Dict[str, str], file_path: str = None
    ):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()

            safe_table = table_name.replace('"', '""')

            columns = []
            placeholders = []
            vals = []

            for col, val in values.items():
                safe_col = col.replace('"', '""')
                columns.append(f'"{safe_col}"')
                placeholders.append("?")
                vals.append(val)

            query = f'INSERT INTO "{safe_table}" ({", ".join(columns)}) VALUES ({", ".join(placeholders)})'
            cursor.execute(query, vals)
            conn.commit()
            new_id = cursor.lastrowid

            return {"success": True, "message": "Ligne ajoutee avec succes.", "row_id": new_id}
        except Exception as e:
            return {"success": False, "message": str(e)}
        finally:
            self._safe_close_connection(conn)

    # ============================================================
    # STATISTIQUES
    # ============================================================
    def get_table_statistics(self, table_name: str, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "stats": {}}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()

            cursor.execute(f'PRAGMA table_info("{table_name}")')
            columns_info = cursor.fetchall()
            columns = [col[1] for col in columns_info]

            stats = {}

            for col in columns:
                col_type = None
                for info in columns_info:
                    if info[1] == col:
                        col_type = info[2].upper()
                        break

                if col_type in ["INTEGER", "REAL", "FLOAT", "NUMERIC"]:
                    cursor.execute(
                        f'SELECT COUNT("{col}"), AVG("{col}"), MIN("{col}"), MAX("{col}"), SUM("{col}") '
                        f'FROM "{table_name}" WHERE "{col}" IS NOT NULL'
                    )
                    result = cursor.fetchone()
                    stats[col] = {
                        "type": "numerique",
                        "count": result[0] if result[0] is not None else 0,
                        "average": round(result[1], 2) if result[1] is not None else None,
                        "min": result[2] if result[2] is not None else None,
                        "max": result[3] if result[3] is not None else None,
                        "sum": round(result[4], 2) if result[4] is not None else None,
                    }
                else:
                    cursor.execute(
                        f'SELECT COUNT(DISTINCT "{col}"), COUNT("{col}") '
                        f'FROM "{table_name}" WHERE "{col}" IS NOT NULL'
                    )
                    distinct, total = cursor.fetchone()
                    stats[col] = {
                        "type": "texte",
                        "distinct_count": distinct if distinct is not None else 0,
                        "total_count": total if total is not None else 0,
                    }

            return {"success": True, "stats": stats, "columns": columns}
        except Exception as e:
            return {"success": False, "message": str(e), "stats": {}}
        finally:
            self._safe_close_connection(conn)

    def get_column_values(self, table_name: str, column: str, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "values": []}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()

            cursor.execute(
                f'SELECT "{column}" FROM "{table_name}" WHERE "{column}" IS NOT NULL AND "{column}" != ""'
            )
            values = [row[0] for row in cursor.fetchall()]

            return {"success": True, "values": values}
        except Exception as e:
            return {"success": False, "message": str(e), "values": []}
        finally:
            self._safe_close_connection(conn)

    def get_table_distribution(self, table_name: str, column: str, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {
                    "success": False,
                    "message": "Aucune base active.",
                    "distribution": [],
                }

            conn = self._connect_db(db_path)
            cursor = conn.cursor()

            cursor.execute(
                f'SELECT "{column}", COUNT(*) FROM "{table_name}" '
                f'WHERE "{column}" IS NOT NULL GROUP BY "{column}" '
                f'ORDER BY COUNT(*) DESC LIMIT 50'
            )
            distribution = [
                {"value": row[0], "count": row[1]} for row in cursor.fetchall()
            ]

            return {"success": True, "distribution": distribution}
        except Exception as e:
            return {"success": False, "message": str(e), "distribution": []}
        finally:
            self._safe_close_connection(conn)

    # ============================================================
    # DOUBLONS  (rowid robuste)
    # ============================================================
    def scan_table_duplicates_advanced(
        self, table_name: str, algorithm: str = "general", file_path: str = None
    ):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "duplicates": []}

            conn = self._connect_db(db_path)
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()

            cursor.execute(f'PRAGMA table_info("{table_name}")')
            columns = [col[1] for col in cursor.fetchall()]
            if not columns:
                return {"success": True, "duplicates": []}

            if algorithm == "general":
                cols_to_check = [c for c in columns if c.lower() != "id"]
                if not cols_to_check:
                    cols_to_check = columns

                df = pd.read_sql_query(
                    f'SELECT rowid AS __sqlite_rowid__, * FROM "{table_name}"',
                    conn,
                )

                if df.empty:
                    return {"success": True, "duplicates": []}

                df["_hash_key"] = (
                    df[cols_to_check]
                    .fillna("")
                    .astype(str)
                    .apply(lambda x: "|".join(x.str.upper().str.strip()), axis=1)
                )

                duplicate_mask = df.duplicated(subset=["_hash_key"], keep="first")
                duplicates = []

                if duplicate_mask.any():
                    grouped = df.groupby("_hash_key")
                    hash_to_first = {
                        hash_val: group.iloc[0] for hash_val, group in grouped
                    }

                    for idx, row in df[duplicate_mask].iterrows():
                        hash_key = row["_hash_key"]
                        ref_row = hash_to_first.get(hash_key)

                        if ref_row is not None:
                            row_dict = self._convert_row_to_dict(row, columns)
                            ref_dict = self._convert_row_to_dict(ref_row, columns)

                            duplicates.append({
                                "row_index": int(row["__sqlite_rowid__"]),
                                "reference_id": int(ref_row["__sqlite_rowid__"]),
                                "data": row_dict,
                                "reference_data": ref_dict,
                                "algorithm": "general",
                            })

                return {"success": True, "duplicates": duplicates, "algorithm": "general"}

            elif algorithm == "cin_nom":
                # Lecture explicite du rowid avec alias unique
                df = pd.read_sql_query(
                    f'SELECT rowid AS __sqlite_rowid__, * FROM "{table_name}"',
                    conn,
                )

                if df.empty:
                    return {"success": True, "duplicates": []}

                cin_col = None
                nom_col = None
                commune_col = None
                fkt_col = None

                cin_patterns = ["cin", "nin", "nif", "id_personne", "num", "numero", "matricule"]
                nom_patterns = ["nom", "name", "prenom", "firstname", "lastname", "fullname", "raison"]
                commune_patterns = ["commune", "comm", "ville", "city"]
                fkt_patterns = ["fkt", "fokontany", "localite", "lieu", "village"]

                for col in columns:
                    col_lower = col.lower()
                    if not cin_col and any(p in col_lower for p in cin_patterns):
                        cin_col = col
                    if not nom_col and any(p in col_lower for p in nom_patterns):
                        nom_col = col
                    if not commune_col and any(p in col_lower for p in commune_patterns):
                        commune_col = col
                    if not fkt_col and any(p in col_lower for p in fkt_patterns):
                        fkt_col = col

                if not cin_col or not nom_col:
                    return {
                        "success": True,
                        "duplicates": [],
                        "algorithm": "cin_nom",
                        "message": "Colonnes CIN ou NOM non trouvees",
                    }

                if not commune_col:
                    return {
                        "success": True,
                        "duplicates": [],
                        "algorithm": "cin_nom",
                        "message": "Colonne COMMUNE non trouvee",
                    }

                if not fkt_col:
                    return {
                        "success": True,
                        "duplicates": [],
                        "algorithm": "cin_nom",
                        "message": "Colonne FKT non trouvee",
                    }

                df["_cin_clean"] = df[cin_col].fillna("").astype(str).str.upper().str.strip()
                df["_cin_clean"] = df["_cin_clean"].str.replace(r"[^A-Z0-9]", "", regex=True)
                df["_nom_clean"] = df[nom_col].fillna("").astype(str).str.upper().str.strip()
                df["_commune_clean"] = (
                    df[commune_col].fillna("").astype(str).str.upper().str.strip()
                )
                df["_fkt_clean"] = df[fkt_col].fillna("").astype(str).str.upper().str.strip()

                df = df[(df["_cin_clean"] != "") | (df["_nom_clean"] != "")].copy()

                if df.empty:
                    return {"success": True, "duplicates": []}

                duplicates = []
                processed = set()

                # Construire rows_list en utilisant __sqlite_rowid__
                rows_list = []
                for idx, row in df.iterrows():
                    row_data = {
                        "rowid": int(row["__sqlite_rowid__"]),
                        "cin": row["_cin_clean"],
                        "nom": row["_nom_clean"],
                        "commune": row["_commune_clean"],
                        "fkt": row["_fkt_clean"],
                        "data": {k: row[k] for k in columns if k in row.index},
                    }
                    rows_list.append(row_data)

                composite_groups = {}
                for row in rows_list:
                    if (
                        row["cin"]
                        and len(row["cin"]) >= 3
                        and row["nom"]
                        and row["commune"]
                        and row["fkt"]
                    ):
                        key = f"{row['cin']}|{row['nom']}|{row['commune']}|{row['fkt']}"

                        if key not in composite_groups:
                            composite_groups[key] = []
                        composite_groups[key].append(row)

                for composite_key, group in composite_groups.items():
                    if len(group) <= 1:
                        continue

                    ref_row = group[0]
                    ref_id = ref_row["rowid"]

                    for row in group[1:]:
                        if row["rowid"] in processed:
                            continue

                        parts = composite_key.split("|")

                        duplicates.append({
                            "row_index": row["rowid"],
                            "reference_id": ref_id,
                            "data": row["data"],
                            "reference_data": ref_row["data"],
                            "algorithm": "cin_nom_complet",
                            "cin_col": cin_col,
                            "nom_col": nom_col,
                            "commune_col": commune_col,
                            "fkt_col": fkt_col,
                            "cin_value": row["cin"],
                            "nom_value": row["nom"],
                            "commune_value": row["commune"],
                            "fkt_value": row["fkt"],
                            "context": {
                                "type": "CIN_NOM_COMMUNE_FKT",
                                "cin": parts[0] if len(parts) > 0 else "",
                                "nom": parts[1] if len(parts) > 1 else "",
                                "commune": parts[2] if len(parts) > 2 else "",
                                "fkt": parts[3] if len(parts) > 3 else "",
                            },
                        })
                        processed.add(row["rowid"])

                return {"success": True, "duplicates": duplicates, "algorithm": "cin_nom"}

            else:
                return {"success": True, "duplicates": []}

        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"success": False, "message": str(e), "duplicates": []}
        finally:
            self._safe_close_connection(conn)

    def _convert_row_to_dict(self, row, columns):
        result = {}
        for k in columns:
            val = row.get(k)
            if val is None or pd.isna(val):
                result[k] = None
            elif hasattr(val, "item"):
                try:
                    result[k] = val.item()
                except Exception:
                    result[k] = str(val)
            elif isinstance(val, (int, float, str, bool)):
                result[k] = val
            else:
                result[k] = str(val)
        return result

    def delete_table_row(self, table_name: str, row_id: int, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()
            cursor.execute(f'DELETE FROM "{table_name}" WHERE rowid = ?', (row_id,))
            conn.commit()
            return {"success": True, "message": f"Ligne {row_id} supprimee."}
        except Exception as e:
            return {"success": False, "message": str(e)}
        finally:
            self._safe_close_connection(conn)

    def clean_database_values(self, file_path: str = None):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active."}

            conn = self._connect_db(db_path)
            cursor = conn.cursor()
            cursor.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
            )
            tables = [
                row[0] for row in cursor.fetchall() if row[0] != "sqlite_sequence"
            ]

            for table in tables:
                df = pd.read_sql_query(f"SELECT * FROM [{table}]", conn)
                if df.empty:
                    continue
                for col in df.columns:
                    if pd.api.types.is_numeric_dtype(df[col]):
                        df[col] = df[col].fillna(0)
                    else:
                        df[col] = df[col].fillna("Non specifie")
                df.to_sql(table, conn, if_exists="replace", index=False)

            return {"success": True, "message": "Nettoyage des valeurs NaN/Null termine."}
        except Exception as e:
            return {"success": False, "message": str(e)}
        finally:
            self._safe_close_connection(conn)

    # ============================================================
    # IMPORTATION EXCEL
    # ============================================================
    def select_excel_file(self):
        global _APP_WINDOW
        if _APP_WINDOW is None:
            return {"success": False, "message": "La fenetre n'est pas disponible."}
        try:
            result = _APP_WINDOW.create_file_dialog(
                webview.OPEN_DIALOG,
                file_types=(
                    "Excel Files (*.xls;*.xlsx)",
                    "Old Excel (*.xls)",
                    "Modern Excel (*.xlsx)",
                ),
            )
            if not result:
                return {"success": False, "message": "Aucun fichier selectionne."}
            return {"success": True, "file_path": str(result[0])}
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de selectionner le fichier : {e}",
            }

    def select_excel_save_file(self):
        global _APP_WINDOW
        if _APP_WINDOW is None:
            return {"success": False, "message": "La fenetre n'est pas disponible."}
        try:
            result = _APP_WINDOW.create_file_dialog(
                webview.SAVE_DIALOG,
                file_types=("Excel Files (*.xlsx)",),
                directory=os.path.expanduser("~/Documents"),
            )

            if not result:
                return {"success": False, "message": "Aucun fichier selectionne."}

            if isinstance(result, (list, tuple)):
                file_path = str(result[0])
            else:
                file_path = str(result)

            if not file_path.lower().endswith(".xlsx"):
                file_path += ".xlsx"

            return {"success": True, "file_path": file_path}
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de selectionner le fichier : {e}",
            }

    def get_excel_sheets(self, file_path: str):
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le chemin du fichier est requis."}
            sheets = self._excel_service.get_sheets(file_path)
            return {"success": True, "file_path": str(file_path), "sheets": sheets or []}
        except Exception as e:
            return {
                "success": False,
                "message": f"Le fichier Excel est invalide : {e}",
                "sheets": [],
            }

    def preview_excel_sheet(self, file_path: str, sheet_name: str):
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le fichier Excel est requis."}
            if not sheet_name or not str(sheet_name).strip():
                return {"success": False, "message": "La feuille Excel est requise."}

            preview = self._excel_service.preview_sheet(
                file_path, sheet_name, max_rows=1
            )
            headers = []
            if isinstance(preview, dict):
                headers = preview.get("headers", [])
                if not headers and "data" in preview and len(preview["data"]) > 0:
                    headers = preview["data"][0]
            return {
                "success": True,
                "headers": headers if isinstance(headers, list) else list(headers),
                "preview": preview.get("data", []) if isinstance(preview, dict) else [],
            }
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de lire la feuille : {e}",
            }

    def _clean_ascii(self, text) -> str:
        nfkd_form = unicodedata.normalize("NFKD", str(text))
        only_ascii = "".join([c for c in nfkd_form if not unicodedata.combining(c)])
        return re.sub(r"[^\w]", "_", only_ascii).lower().strip("_")

    def import_excel_to_database(
        self, file_path: str, sheet_name: str = None, table_name: str = None
    ):
        """
        Import renforce : ferme toute base ouverte avant, libere les ressources
        apres l'ecriture, puis ouvre la nouvelle base sans conflit de thread.
        """
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le fichier Excel est requis."}

            excel_path = Path(file_path)
            data_dir = DATA_DIR
            data_dir.mkdir(parents=True, exist_ok=True)

            safe_db_name = self._clean_ascii(excel_path.stem) or "database"
            db_filename = f"{safe_db_name}.db"
            db_path = data_dir / db_filename

            # Fermer toute base ouverte AVANT l'import
            if self._active_db_path:
                try:
                    self._database_service.close_database()
                except Exception:
                    pass
                self._active_db_path = None
                self._last_db_path = None
                self._release_resources(0.2)

            all_sheets = pd.read_excel(excel_path, sheet_name=None)
            conn = self._connect_db(str(db_path))
            tables_created = []
            total_rows = 0
            log_path = data_dir / "logs.txt"

            try:
                for current_sheet, df in all_sheets.items():
                    if sheet_name and str(current_sheet) != str(sheet_name):
                        continue

                    clean_table_name = (
                        table_name
                        if (table_name and len(all_sheets) == 1)
                        else (self._clean_ascii(str(current_sheet)) or "table")
                    )

                    try:
                        df = df.dropna(how="all")

                        if any(str(col).lower().startswith("unnamed") for col in df.columns):
                            if len(df) > 0:
                                new_headers = (
                                    df.iloc[0].fillna("colonne_inconnue").astype(str).tolist()
                                )
                                cleaned_headers = []
                                seen = {}
                                for h in new_headers:
                                    h_clean = self._clean_ascii(h) or "col"
                                    if h_clean in seen:
                                        seen[h_clean] += 1
                                        h_clean = f"{h_clean}_{seen[h_clean]}"
                                    else:
                                        seen[h_clean] = 0
                                    cleaned_headers.append(h_clean)
                                df.columns = cleaned_headers
                                df = df.drop(df.index[0])

                        for col in df.columns:
                            if pd.api.types.is_numeric_dtype(df[col]):
                                df[col] = df[col].fillna(0)
                            elif pd.api.types.is_datetime64_any_dtype(df[col]):
                                df[col] = pd.to_datetime(df[col]).dt.date
                                df[col] = df[col].fillna(pd.Timestamp.now().date())
                            else:
                                df[col] = df[col].fillna("Non specifie")

                        df.columns = [self._clean_ascii(str(col)) or "col" for col in df.columns]
                        df.to_sql(clean_table_name, conn, if_exists="replace", index=False)
                        tables_created.append(clean_table_name)
                        total_rows += len(df)
                    except Exception as sheet_err:
                        err_msg = f"[ERREUR IMPORT] Feuille '{current_sheet}' : {str(sheet_err)}\n"
                        with open(log_path, "a", encoding="utf-8") as log_file:
                            log_file.write(err_msg)
                        continue
            finally:
                self._safe_close_connection(conn)

            # Liberer les ressources avant d'ouvrir la nouvelle base
            self._release_resources(0.35)

            open_result = self.open_database(str(db_path))

            return {
                "success": True,
                "message": f"Importation reussie ! {len(tables_created)} table(s) creee(s).",
                "tables": tables_created,
                "total_rows": total_rows,
                "db_path": str(db_path),
                "open_result": open_result,
            }
        except Exception as e:
            import traceback
            traceback.print_exc()
            return {"success": False, "message": f"Erreur lors de la conversion : {e}"}

    # ============================================================
    # EXPORTATION PDF
    # ============================================================
    def select_pdf_file(self):
        global _APP_WINDOW
        if _APP_WINDOW is None:
            return {"success": False, "message": "La fenetre n'est pas disponible."}
        try:
            result = _APP_WINDOW.create_file_dialog(
                webview.SAVE_DIALOG,
                file_types=("PDF Files (*.pdf)",),
                directory=os.path.expanduser("~/Documents"),
            )

            if not result:
                return {"success": False, "message": "Aucun fichier selectionne."}

            if isinstance(result, (list, tuple)):
                file_path = str(result[0])
            else:
                file_path = str(result)

            if not file_path.lower().endswith(".pdf"):
                file_path += ".pdf"

            return {"success": True, "file_path": file_path}
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de selectionner le fichier PDF : {e}",
            }

    def generate_pdf_from_html(self, output_path: str, html_content: str):
        try:
            output_path = str(output_path).strip()

            if output_path.startswith("('") or output_path.startswith('("'):
                import ast

                try:
                    parsed = ast.literal_eval(output_path)
                    if isinstance(parsed, (list, tuple)):
                        output_path = str(parsed[0])
                except Exception:
                    output_path = (
                        output_path.replace("('", "")
                        .replace("')", "")
                        .replace('("', "")
                        .replace('")', "")
                    )

            output_dir = os.path.dirname(output_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir, exist_ok=True)

            try:
                from weasyprint import HTML

                HTML(string=html_content).write_pdf(output_path)
                return {"success": True, "message": f"PDF genere avec succes : {output_path}"}
            except ImportError:
                pass

            try:
                import pdfkit

                options = {
                    "page-size": "A4",
                    "margin-top": "0.5in",
                    "margin-right": "0.5in",
                    "margin-bottom": "0.5in",
                    "margin-left": "0.5in",
                    "encoding": "UTF-8",
                    "no-outline": None,
                    "enable-local-file-access": None,
                }
                pdfkit.from_string(html_content, output_path, options=options)
                return {"success": True, "message": f"PDF genere avec succes : {output_path}"}
            except ImportError:
                html_path = output_path.replace(".pdf", ".html")
                with open(html_path, "w", encoding="utf-8") as f:
                    f.write(html_content)
                return {"success": True, "message": f"HTML genere avec succes : {html_path}"}

        except Exception as e:
            import traceback

            traceback.print_exc()
            return {
                "success": False,
                "message": f"Erreur lors de la generation du PDF : {e}",
            }

    # ============================================================
    # EXPORTATION EXCEL (fichier de sortie)
    # ============================================================
    def select_excel_export_file(self):
        global _APP_WINDOW
        if _APP_WINDOW is None:
            return {"success": False, "message": "La fenetre n'est pas disponible."}
        try:
            result = _APP_WINDOW.create_file_dialog(
                webview.SAVE_DIALOG,
                file_types=("Excel Files (*.xlsx)",),
                directory=os.path.expanduser("~/Documents"),
            )

            if not result:
                return {"success": False, "message": "Aucun fichier selectionne."}

            if isinstance(result, (list, tuple)):
                file_path = str(result[0])
            else:
                file_path = str(result)

            if not file_path.lower().endswith(".xlsx"):
                file_path += ".xlsx"

            return {"success": True, "file_path": file_path}
        except Exception as e:
            return {
                "success": False,
                "message": f"Impossible de selectionner le fichier Excel : {e}",
            }

    def generate_excel_from_data(
        self, output_path: str, data: List[dict], headers: List[str] = None
    ):
        try:
            output_path = str(output_path).strip()

            if output_path.startswith("('") or output_path.startswith('("'):
                import ast

                try:
                    parsed = ast.literal_eval(output_path)
                    if isinstance(parsed, (list, tuple)):
                        output_path = str(parsed[0])
                except Exception:
                    output_path = (
                        output_path.replace("('", "")
                        .replace("')", "")
                        .replace('("', "")
                        .replace('")', "")
                    )

            output_dir = os.path.dirname(output_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir, exist_ok=True)

            df = pd.DataFrame(data)
            if headers:
                if len(headers) == len(df.columns):
                    df.columns = headers

            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name="Resultats", index=False)
                self._apply_professional_formatting(writer, "Resultats", df)

            return {"success": True, "message": f"Excel genere avec succes : {output_path}"}
        except Exception as e:
            return {
                "success": False,
                "message": f"Erreur lors de la generation de l'Excel : {e}",
            }

    # ============================================================
    # REQUETES SQL AVANCEES
    # ============================================================
    def execute_custom_sql_operation(
        self,
        table_name: str,
        op_type: str,
        attribute: str = None,
        value: str = None,
        group_by: str = None,
        file_path: str = None,
    ):
        conn = None
        try:
            db_path = self._get_db_path(file_path)
            if not db_path:
                return {"success": False, "message": "Aucune base active.", "data": []}

            safe_table = table_name.replace('"', '""')
            safe_attr = attribute.replace('"', '""') if attribute else "*"
            safe_group = group_by.replace('"', '""') if group_by else None

            conn = self._connect_db(db_path)
            op = op_type.upper()

            query = ""
            params = ()

            if op == "SELECT_ALL":
                query = f'SELECT * FROM "{safe_table}" LIMIT 1000'

            elif op == "DISTINCT":
                if not attribute:
                    return {
                        "success": False,
                        "message": "Attribut requis pour DISTINCT.",
                        "data": [],
                    }
                query = f'SELECT DISTINCT "{safe_attr}" FROM "{safe_table}"'

            elif op in ["MIN", "MAX", "COUNT", "SUM", "AVG"]:
                if not attribute and op != "COUNT":
                    return {
                        "success": False,
                        "message": f"Attribut requis pour {op}.",
                        "data": [],
                    }
                target_col = f'"{safe_attr}"' if attribute else "*"
                query = f'SELECT {op}({target_col}) as result FROM "{safe_table}"'

            elif op == "WHERE_LIKE":
                if not attribute:
                    return {
                        "success": False,
                        "message": "Attribut requis pour WHERE / LIKE.",
                        "data": [],
                    }
                query = f'SELECT * FROM "{safe_table}" WHERE "{safe_attr}" LIKE ?'
                params = (f"%{value if value else ''}%",)

            elif op == "GROUP_BY":
                if not safe_group:
                    return {
                        "success": False,
                        "message": "Champ de groupement requis.",
                        "data": [],
                    }
                agg_col = f'"{safe_attr}"' if attribute else "*"
                query = (
                    f'SELECT "{safe_group}", COUNT({agg_col}) as total '
                    f'FROM "{safe_table}" GROUP BY "{safe_group}"'
                )

            else:
                query = f'SELECT * FROM "{safe_table}" LIMIT 1000'

            df = pd.read_sql_query(query, conn, params=params if params else None)

            for col in df.columns:
                if pd.api.types.is_datetime64_any_dtype(df[col]):
                    df[col] = df[col].dt.date

            return {"success": True, "data": df.fillna("").to_dict(orient="records")}
        except Exception as e:
            return {
                "success": False,
                "message": f"Erreur SQL : {str(e)}",
                "data": [],
            }
        finally:
            self._safe_close_connection(conn)


def main():
    global _APP_WINDOW

    # Initialiser system.db dans le dossier PERSISTANT
    initialize_database()

    api = Api()

    _APP_WINDOW = webview.create_window(
        "Data Manager - Expert Edition",
        str(INDEX_FILE),
        js_api=api,
        width=1280,
        height=800,
        resizable=True,
        fullscreen=False,
    )
    webview.start()


if __name__ == "__main__":
    main()