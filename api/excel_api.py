"""
Import et export Excel.
"""

import warnings

warnings.filterwarnings(
    "ignore",
    message="Cell .* is marked as a date but the serial value .* is outside",
    category=UserWarning,
    module="openpyxl",
)

import os
import unicodedata
import re
from pathlib import Path
from typing import List, Dict, Any

import pandas as pd

from paths import DATA_DIR, get_app_window
from api.utils import connect_db, safe_close_connection, release_resources, clean_ascii


class ExcelApi:
    def __init__(self, database_api, excel_service, database_service):
        self._database_api = database_api
        self._excel_service = excel_service
        self._database_service = database_service

    # ============================================================
    # SELECTION DE FICHIERS
    # ============================================================
    def select_excel_file(self):
        import webview

        window = get_app_window()
        if window is None:
            return {"success": False, "message": "La fenetre n'est pas disponible."}
        try:
            result = window.create_file_dialog(
                webview.OPEN_DIALOG,
                file_types=(
                    "Excel Files (*.xls;*.xlsx)",
                    "Old Excel (*.xls)",
                    "Modern Excel (*.xlsx)",
                ),
            )
            if not result:
                return {"success": False, "message": "Aucun fichier selectionne."}
            return {"success": True, "file_path": str(result[0])}
        except Exception as e:
            return {"success": False, "message": f"Impossible de selectionner le fichier : {e}"}

    def select_excel_save_file(self):
        import webview

        window = get_app_window()
        if window is None:
            return {"success": False, "message": "La fenetre n'est pas disponible."}
        try:
            result = window.create_file_dialog(
                webview.SAVE_DIALOG,
                file_types=("Excel Files (*.xlsx)",),
                directory=os.path.expanduser("~/Documents"),
            )
            if not result:
                return {"success": False, "message": "Aucun fichier selectionne."}

            file_path = str(result[0]) if isinstance(result, (list, tuple)) else str(result)
            if not file_path.lower().endswith(".xlsx"):
                file_path += ".xlsx"

            return {"success": True, "file_path": file_path}
        except Exception as e:
            return {"success": False, "message": f"Impossible de selectionner le fichier : {e}"}

    def select_excel_export_file(self):
        return self.select_excel_save_file()

    # ============================================================
    # APERCU EXCEL
    # ============================================================
    def get_excel_sheets(self, file_path: str):
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le chemin du fichier est requis."}
            sheets = self._excel_service.get_sheets(file_path)
            return {"success": True, "file_path": str(file_path), "sheets": sheets or []}
        except Exception as e:
            return {"success": False, "message": f"Le fichier Excel est invalide : {e}", "sheets": []}

    def preview_excel_sheet(self, file_path: str, sheet_name: str):
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le fichier Excel est requis."}
            if not sheet_name or not str(sheet_name).strip():
                return {"success": False, "message": "La feuille Excel est requise."}

            preview = self._excel_service.preview_sheet(file_path, sheet_name, max_rows=1)
            headers = []
            if isinstance(preview, dict):
                headers = preview.get("headers", [])
                if not headers and "data" in preview and len(preview["data"]) > 0:
                    headers = preview["data"][0]
            return {
                "success": True,
                "headers": headers if isinstance(headers, list) else list(headers),
                "preview": preview.get("data", []) if isinstance(preview, dict) else [],
            }
        except Exception as e:
            return {"success": False, "message": f"Impossible de lire la feuille : {e}"}

    # ============================================================
    # IMPORT EXCEL -> SQLITE
    # ============================================================
    def import_excel_to_database(self, file_path: str, sheet_name: str = None, table_name: str = None):
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le fichier Excel est requis."}

            excel_path = Path(file_path)
            data_dir = DATA_DIR
            data_dir.mkdir(parents=True, exist_ok=True)

            safe_db_name = clean_ascii(excel_path.stem) or "database"
            db_filename = f"{safe_db_name}.db"
            db_path = data_dir / db_filename

            if self._database_api.active_db_path:
                try:
                    self._database_service.close_database()
                except Exception:
                    pass
                self._database_api.active_db_path = None
                release_resources(0.2)

            all_sheets = pd.read_excel(excel_path, sheet_name=None)
            conn = connect_db(str(db_path))
            tables_created = []
            total_rows = 0
            log_path = data_dir / "logs.txt"

            try:
                for current_sheet, df in all_sheets.items():
                    if sheet_name and str(current_sheet) != str(sheet_name):
                        continue

                    clean_table_name = (
                        table_name
                        if (table_name and len(all_sheets) == 1)
                        else (clean_ascii(str(current_sheet)) or "table")
                    )

                    try:
                        df = df.dropna(how="all")

                        if any(str(col).lower().startswith("unnamed") for col in df.columns):
                            if len(df) > 0:
                                new_headers = df.iloc[0].fillna("colonne_inconnue").astype(str).tolist()
                                cleaned_headers = []
                                seen = {}
                                for h in new_headers:
                                    h_clean = clean_ascii(h) or "col"
                                    if h_clean in seen:
                                        seen[h_clean] += 1
                                        h_clean = f"{h_clean}_{seen[h_clean]}"
                                    else:
                                        seen[h_clean] = 0
                                    cleaned_headers.append(h_clean)
                                df.columns = cleaned_headers
                                df = df.drop(df.index[0])

                        for col in df.columns:
                            if pd.api.types.is_numeric_dtype(df[col]):
                                df[col] = df[col].fillna(0)
                            elif pd.api.types.is_datetime64_any_dtype(df[col]):
                                df[col] = pd.to_datetime(df[col]).dt.date
                                df[col] = df[col].fillna(pd.Timestamp.now().date())
                            else:
                                df[col] = df[col].fillna("Non specifie").infer_objects(copy=False)

                        df.columns = [clean_ascii(str(col)) or "col" for col in df.columns]
                        df.to_sql(clean_table_name, conn, if_exists="replace", index=False)
                        tables_created.append(clean_table_name)
                        total_rows += len(df)
                    except Exception as sheet_err:
                        err_msg = f"[ERREUR IMPORT] Feuille '{current_sheet}' : {str(sheet_err)}\n"
                        with open(log_path, "a", encoding="utf-8") as log_file:
                            log_file.write(err_msg)
                        continue
            finally:
                safe_close_connection(conn)

            release_resources(0.35)
            open_result = self._database_api.open_database(str(db_path))

            return {
                "success": True,
                "message": f"Importation reussie ! {len(tables_created)} table(s) creee(s).",
                "tables": tables_created,
                "total_rows": total_rows,
                "db_path": str(db_path),
                "open_result": open_result,
            }
        except Exception as e:
            return {"success": False, "message": f"Erreur lors de la conversion : {e}"}

    def create_database_from_excel(self, file_path: str, db_name: str = None):
        """Cree une base SQLite a partir d'un fichier Excel (sans ouvrir)."""
        try:
            if not file_path or not str(file_path).strip():
                return {"success": False, "message": "Le fichier Excel est requis."}

            excel_path = Path(file_path)
            if not os.path.exists(str(excel_path)):
                return {"success": False, "message": f"Fichier introuvable : {file_path}"}

            data_dir = DATA_DIR
            data_dir.mkdir(parents=True, exist_ok=True)

            safe_db_name = clean_ascii(db_name or excel_path.stem) or "database"
            db_filename = f"{safe_db_name}.db"
            db_path = data_dir / db_filename

            all_sheets = pd.read_excel(excel_path, sheet_name=None)
            conn = connect_db(str(db_path))
            tables_created = []
            total_rows = 0

            try:
                for current_sheet, df in all_sheets.items():
                    clean_table_name = clean_ascii(str(current_sheet)) or "table"
                    try:
                        df = df.dropna(how="all")
                        for col in df.columns:
                            if pd.api.types.is_numeric_dtype(df[col]):
                                df[col] = df[col].fillna(0)
                            else:
                                df[col] = df[col].fillna("Non specifie")
                        df.columns = [clean_ascii(str(col)) or "col" for col in df.columns]
                        df.to_sql(clean_table_name, conn, if_exists="replace", index=False)
                        tables_created.append(clean_table_name)
                        total_rows += len(df)
                    except Exception as sheet_err:
                        print(f"[ERREUR IMPORT] Feuille '{current_sheet}' : {sheet_err}")
                        continue
            finally:
                safe_close_connection(conn)

            return {
                "success": True,
                "message": f"Base de donnees '{db_filename}' creee avec {len(tables_created)} table(s).",
                "db_path": str(db_path),
                "db_name": db_filename,
                "tables": tables_created,
                "total_rows": total_rows,
            }
        except Exception as e:
            return {"success": False, "message": f"Erreur lors de la creation : {e}"}

    # ============================================================
    # FORMATAGE EXCEL
    # ============================================================
    def _apply_professional_formatting(self, writer, sheet_name: str, df: pd.DataFrame):
        try:
            from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
            from openpyxl.utils import get_column_letter

            worksheet = writer.sheets[sheet_name]
            max_row = len(df) + 1
            max_col = len(df.columns)

            if max_col == 0:
                return

            worksheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

            header_font = Font(bold=True, color="FFFFFF", size=11)
            header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
            header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            thin_border = Border(
                left=Side(style="thin", color="E2E8F0"),
                right=Side(style="thin", color="E2E8F0"),
                top=Side(style="thin", color="E2E8F0"),
                bottom=Side(style="thin", color="E2E8F0"),
            )

            for col_idx in range(1, max_col + 1):
                cell = worksheet.cell(row=1, column=col_idx)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
                cell.border = thin_border

            alt_fill_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
            alt_fill_2 = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            cell_alignment = Alignment(vertical="center", wrap_text=False)

            for row_idx in range(2, max_row + 1):
                fill = alt_fill_1 if row_idx % 2 == 0 else alt_fill_2
                for col_idx in range(1, max_col + 1):
                    cell = worksheet.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    cell.fill = fill
                    cell.alignment = cell_alignment

            worksheet.freeze_panes = "A2"

            for col_idx, column in enumerate(df.columns, 1):
                col_letter = get_column_letter(col_idx)
                max_length = len(str(column))
                for row_idx in range(2, min(len(df) + 2, 500)):
                    cell_value = worksheet.cell(row=row_idx, column=col_idx).value
                    if cell_value is not None:
                        cell_len = len(str(cell_value))
                        if cell_len > max_length:
                            max_length = cell_len
                adjusted_width = min(max_length + 3, 50)
                worksheet.column_dimensions[col_letter].width = adjusted_width

            worksheet.row_dimensions[1].height = 25
        except Exception as e:
            print(f"[WARN] Erreur formatage Excel pour '{sheet_name}' : {e}")

    # ============================================================
    # EXPORT SQLITE -> EXCEL
    # ============================================================
    def export_database_to_excel_from_path(self, db_path: str, output_excel_path: str):
        """
        Export professionnel d'une base SQLite vers Excel.
        - Chaque table devient une feuille
        - Les attributs deviennent les en-tetes
        - AutoFilter + mise en forme
        """
        try:
            if not db_path or not os.path.exists(db_path):
                return {"success": False, "message": "Base de donnees introuvable."}

            if not output_excel_path or not str(output_excel_path).strip():
                return {"success": False, "message": "Le chemin de sortie est requis."}

            output_dir = os.path.dirname(output_excel_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir, exist_ok=True)

            conn = connect_db(db_path)
            try:
                cursor = conn.cursor()
                cursor.execute(
                    "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';"
                )
                tables = [row[0] for row in cursor.fetchall() if row[0] != "sqlite_sequence"]

                if not tables:
                    return {"success": False, "message": "Aucune table dans cette base."}

                tables_ordered = [t for t in tables if t != "listes_meres"]
                if "listes_meres" in tables:
                    tables_ordered.append("listes_meres")

                tables_exported = []
                total_rows_exported = 0

                with pd.ExcelWriter(output_excel_path, engine="openpyxl") as writer:
                    for table in tables_ordered:
                        df = pd.read_sql_query(f'SELECT * FROM "{table}"', conn)
                        if "id" in df.columns:
                            df = df.drop(columns=["id"])
                        for col in df.columns:
                            if pd.api.types.is_datetime64_any_dtype(df[col]):
                                df[col] = df[col].dt.date

                        sheet_name = table[:31] if len(table) > 31 else table
                        df.to_excel(writer, sheet_name=sheet_name, index=False)
                        self._apply_professional_formatting(writer, sheet_name, df)
                        tables_exported.append(sheet_name)
                        total_rows_exported += len(df)

                return {
                    "success": True,
                    "message": f"Exportation reussie : {len(tables_exported)} feuille(s) ({total_rows_exported} lignes).",
                    "tables_exported": tables_exported,
                    "total_rows": total_rows_exported,
                    "output_path": output_excel_path,
                }
            finally:
                safe_close_connection(conn)
        except Exception as e:
            return {"success": False, "message": f"Erreur lors de l'export Excel : {e}"}

    def export_database_to_excel(self, output_excel_path: str, file_path: str = None):
        db_path = self._database_api._get_db_path(file_path)
        if not db_path:
            return {"success": False, "message": "Aucune base active."}
        return self.export_database_to_excel_from_path(db_path, output_excel_path)

    def generate_excel_from_data(self, output_path: str, data: List[dict], headers: List[str] = None):
        try:
            output_path = str(output_path).strip()
            output_dir = os.path.dirname(output_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir, exist_ok=True)

            df = pd.DataFrame(data)
            if headers and len(headers) == len(df.columns):
                df.columns = headers

            with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
                df.to_excel(writer, sheet_name="Resultats", index=False)
                self._apply_professional_formatting(writer, "Resultats", df)

            return {"success": True, "message": f"Excel genere avec succes : {output_path}"}
        except Exception as e:
            return {"success": False, "message": f"Erreur lors de la generation de l'Excel : {e}"}